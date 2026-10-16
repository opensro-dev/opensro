# ===========================================================================
#
# pet_design_export.py - the attack-pet design workbook as the game's data
#
# docs/pet-expansion/pet-design.xlsx is the single source of the attack-pet
# expansion (port-only, SRO_ATTACK_PET_EXPANSION): rules, skills, equipment,
# degrees, affixes and procs. This tool reads its saved values (Excel and
# LibreOffice store every formula's result) and writes one JSON document for
# the server (go:embed) and one for the browser client. Edit the workbook,
# then run:
#
#     python scripts/tools/pet_design_export.py
#
# It refuses a workbook with missing values, duplicate IDs or unknown
# names, so a typo never ships a different game.
#
# ===========================================================================
import json
import pathlib
import sys

from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "docs" / "pet-expansion" / "pet-design.xlsx"
OUTPUTS = [
	ROOT / "apps" / "server" / "internal" / "game" / "companion" / "petdesign.json",
	ROOT / "scripts" / "build" / "shared" / "pet-design.json",
]
# The browser client reads the same document as a generated TypeScript module.
CLIENT_MODULE = ROOT / "apps" / "client-next" / "src" / "engine" / "foundation" / "ui" / "pet-design-data.ts"
ROWS = ["Active 1", "Active 2", "Passive 1", "Passive 2", "Passive 3"]
PET_PREFIX = {"Bear": "COS_P_BEAR", "Wolf": "COS_P_WOLF", "Penguin": "COS_P_PENGUIN", "Crow": "COS_P_RAVEN"}
SLOTS = ["Collar", "Harness", "Charm"]


# ================
# fail
# ================
def fail(message):
	sys.exit("pet_design_export: " + message)


# ================
# table
#
# The rows under a sheet's header row 4 as dicts keyed by header text,
# stopping at the first row whose first cell is empty.
# ================
def table(ws, first_column=1):
	names = [c.value for c in ws[4]]
	rows = []
	for r in range(5, ws.max_row + 1):
		values = [ws.cell(row=r, column=c).value for c in range(1, len(names) + 1)]
		if values[first_column - 1] in (None, ""):
			break
		rows.append({n: v for n, v in zip(names, values) if n})
	return rows


# ================
# number
# ================
def number(row, key, where, allow_blank=False):
	value = row.get(key)
	if value in (None, ""):
		if allow_blank:
			return None
		fail(f"{where}: '{key}' is empty (open and save the workbook in Excel or LibreOffice to compute it)")
	if not isinstance(value, (int, float)):
		fail(f"{where}: '{key}' is not a number: {value!r}")
	return int(value) if float(value).is_integer() else float(value)


# ================
# rules
# ================
def rules(ws):
	raw = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(5, 29)}
	get = lambda k: number({"v": raw.get(k)}, "v", "Rules " + k)
	return {
		"skillMaxLevel": get("Skill max level"),
		"rank2Level": get("Rank II level"),
		"rank3Level": get("Rank III level"),
		"petLevelsPerSkillLevel": get("Pet levels per skill level"),
		"pointsPerPetLevels": get("Points per pet levels"),
		"optionUnlockStep": get("Option unlock step"),
		"rowUnlock": [get("Row unlock: " + row) for row in ROWS],
		"passivesPerRow": get("Passives slotted per row"),
		"procInternalCooldownMs": int(get("Proc internal cooldown") * 1000),
		"maxBlues": get("Max blues per item"),
		"magicTwoBluesChance": get("Magic: chance of 2 blues"),
		"rankMultipliers": {k.lower(): get("Monster rank x " + k) for k in ("General", "Champion", "Giant", "Unique")},
		"implicitPerDegree": {s.lower(): get(s + " implicit per degree") for s in SLOTS},
	}


# ================
# main
# ================
def main():
	wb = load_workbook(WORKBOOK, data_only=True)
	design = {
		"generatedBy": "scripts/tools/pet_design_export.py from docs/pet-expansion/pet-design.xlsx - do not edit",
		"rules": rules(wb["Rules"]),
		"pets": [],
		"skills": [],
		"gear": [],
		"affixes": [],
		"procs": [],
	}
	for p in table(wb["Pets"]):
		if p["Pet"] not in PET_PREFIX:
			fail(f"Pets: unknown pet {p['Pet']!r}")
		design["pets"].append({"name": p["Pet"], "role": p["Role"], "codenamePrefix": PET_PREFIX[p["Pet"]]})
	role_names = set()
	for row in wb["Combat roles"].iter_rows(min_row=2, max_row=5, values_only=True):
		name, hp, physical, magical, defense, magic_defense, magic_primary = row[:7]
		if name not in PET_PREFIX or name in role_names or any(not isinstance(v, (int, float)) or not 1 <= v <= 500 for v in (hp, physical, magical, defense, magic_defense)) or magic_primary not in (0, 1):
			fail(f"Combat roles: invalid profile {name!r}")
		role_names.add(name)
		pet = next(p for p in design["pets"] if p["name"] == name)
		pet.update(hpPercent=hp, physicalPercent=physical, magicalPercent=magical, defensePercent=defense, magicDefensePercent=magic_defense, magicPrimary=bool(magic_primary))
	ids = set()
	for s in table(wb["Skills"]):
		where = f"Skills {s['Pet']} {s['Skill']}"
		sid = number(s, "ID", where)
		if sid in ids or not 1 <= sid <= 255:
			fail(f"{where}: ID {sid} is a duplicate or outside 1..255")
		ids.add(sid)
		if s["Pet"] not in PET_PREFIX or s["Row"] not in ROWS:
			fail(f"{where}: unknown pet or row {s['Row']!r}")
		design["skills"].append({
			"id": sid,
			"pet": s["Pet"],
			"row": ROWS.index(s["Row"]),
			"active": s["Row"].startswith("Active"),
			"theme": s["Row theme"],
			"option": number(s, "Option", where),
			"name": s["Skill"],
			"effect": s["Effect ({v})"],
			"trigger": s["Trigger / cooldown kind"],
			"cooldown": number(s, "Cooldown L1", where, True),
			"cooldownUnit": s["Cooldown unit"] or "",
			"cooldownPerLevel": number(s, "Cooldown per level", where, True),
			"value": number(s, "Value L1", where),
			"valuePerLevel": number(s, "Value per level", where),
			"valueUnit": s["Value unit"] or "",
			"maxLevel": number(s, "Max level", where),
			"unlockLevel": number(s, "Unlock pet level", where),
			"rank2": s["Rank II"],
			"rank3": s["Rank III"],
			"animation": s["Animation"] or "",
			"status": s["Status"],
			"icon": s["Icon"] + ".ddj",
		})
	# The approved sequential tree references preserved effect templates.
	design["treeSkills"] = []
	by_id = {s["id"]: s for s in design["skills"]}
	for pet, row, template, first, second, third, last in wb["Skill families"].iter_rows(min_row=2, max_row=21, max_col=7, values_only=True):
		base = by_id.get(template)
		if base is None or base["pet"] != pet or base["row"] != row or not 0 < first < second < third < last <= 255:
			fail(f"Skill families: invalid family {pet} row {row}")
		first_id = min(s["id"] for s in design["skills"] if s["pet"] == pet and s["row"] == row)
		for tier, (start, end) in enumerate(zip((first, second, third), (second, third, last)), 1):
			s = dict(base)
			s.update(id=first_id+tier-1, option=tier, name=base["name"]+" "+("I", "II", "III")[tier-1],
				family=base["name"], unlockLevel=start, tierEnd=end,
				value=base["value"]+base["valuePerLevel"]*base["maxLevel"]*(tier-1))
			if s["family"] == "Challenging Roar":
				s.update(status="Implemented", value=3 if tier < 3 else 6, valuePerLevel=0,
					effect="Taunts up to {v} monsters. Radius grows with proficiency; party monsters require Tier II. No damage.")
			design["treeSkills"].append(s)
	degrees = {int(d["Degree"]): d for d in table(wb["Degrees"])}
	for g in table(wb["Gear"]):
		d = int(g["Degree"])
		dd = degrees.get(d) or fail(f"Gear: degree {d} missing on Degrees")
		where = f"Degrees D{d}"
		design["gear"].append({
			"degree": d,
			"minPetLevel": number(g, "Min pet level", f"Gear D{d}"),
			"dropChance": number(dd, "Drop % per kill (General)", where),
			"magicChance": number(dd, "Magic chance", where),
			"rareChance": number(dd, "Rare chance", where),
			"rareWeights": [number(dd, f"Rare: {n} blues weight", where) for n in (3, 4, 5, 6)],
			"maxBlues": number(dd, "Max blues", where),
			"rollFloor": number(dd, "Minimum roll percentile", where),
			"items": [{
				"slot": slot,
				"name": g[slot + " base"],
				"codename": g[slot + " codename"],
				"id": number(g, slot + " ID", f"Gear D{d}"),
				# Port-only balance: flat stats scaled to the degree's level band.
				"baseStats": {
					"physicalAttack": (int(g["Min pet level"]) + 2) * (2 if k == 0 else 1 if k == 2 else 0),
					"magicalAttack": (int(g["Min pet level"]) + 2) * (2 if k == 0 else 1 if k == 2 else 0),
					"physicalDefense": (int(g["Min pet level"]) + 2) // (1 if k == 1 else 2) if k != 0 else 0,
					"magicalDefense": (int(g["Min pet level"]) + 2) // (1 if k == 1 else 2) if k != 0 else 0,
					"hitRatio": d * (2 if k == 0 else 1 if k == 2 else 0),
					"parryRatio": d * (2 if k == 1 else 1 if k == 2 else 0),
				},
			} for k, slot in enumerate(SLOTS)],
		})
	affix_ids = set()
	for a in table(wb["Affixes"]):
		where = f"Affixes {a['Affix ID']}"
		if a["Affix ID"] in affix_ids:
			fail(f"{where}: duplicate affix ID")
		affix_ids.add(a["Affix ID"])
		if a["Kind"] not in ("Prefix", "Suffix"):
			fail(f"{where}: Kind must be Prefix or Suffix")
		slots = [s.strip() for s in str(a["Slots"]).split(",")]
		if not set(slots) <= set(SLOTS):
			fail(f"{where}: unknown slot in {a['Slots']!r}")
		design["affixes"].append({
			"id": a["Affix ID"], "name": a["Name"], "kind": a["Kind"], "group": a["Group"], "category": a["Category"],
			"slots": slots, "effect": a["Effect ({v})"], "tier": number(a, "Tier", where),
			"minDegree": number(a, "Min degree", where), "min": number(a, "Value min", where),
			"max": number(a, "Value max", where), "unit": a["Unit"], "weight": number(a, "Weight", where),
			"pets": a["Pets"], "status": a["Status"], "code": number(a, "Code", where),
		})
	for p in table(wb["Procs"]):
		design["procs"].append({"name": p["Proc"], "effect": p["Effect ({v})"], "base": number(p, "Value at level 1", p["Proc"]),
			"perLevel": number(p, "Value per level", p["Proc"]), "unit": p["Unit"]})
	text = json.dumps(design, ensure_ascii=False, indent="\t") + "\n"
	for out in OUTPUTS:
		out.write_text(text, encoding="utf-8", newline="\n")
		print("wrote", out.relative_to(ROOT), len(design["skills"]), "skills", len(design["affixes"]), "affixes")
	module = ("// Generated by scripts/tools/pet_design_export.py from docs/pet-expansion/pet-design.xlsx - do not edit.\n"
		+ "export const PET_DESIGN = " + text.rstrip("\n") + " as const;\n")
	CLIENT_MODULE.write_text(module, encoding="utf-8", newline="\n")
	print("wrote", CLIENT_MODULE.relative_to(ROOT))


if __name__ == "__main__":
	main()
